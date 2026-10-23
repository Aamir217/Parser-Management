import hashlib
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)
from openpyxl.utils import get_column_letter

from app.core.config import settings


class ValidatedXlsxError(ValueError):
    def __init__(
        self,
        message: str,
        guardrail_name: str,
    ) -> None:
        super().__init__(message)

        self.message = message
        self.guardrail_name = guardrail_name


def excel_safe_value(
    value: Any,
) -> Any:
    """
    Convert a Python value into a safe value that can
    be written to an Excel cell.

    Formula-like strings are prefixed with an
    apostrophe to prevent spreadsheet formula
    injection when users open the downloaded file.
    """

    if value is None:
        return None

    if isinstance(
        value,
        (
            bool,
            int,
            float,
            datetime,
            date,
        ),
    ):
        return value

    if isinstance(
        value,
        (
            dict,
            list,
            tuple,
            set,
        ),
    ):
        value = json.dumps(
            value,
            ensure_ascii=False,
            default=str,
            sort_keys=True,
        )

    text = str(value)

    stripped_text = text.lstrip()

    if stripped_text.startswith(
        (
            "=",
            "+",
            "-",
            "@",
        )
    ):
        return "'" + text

    return text


def sanitize_output_filename_stem(
    filename: str,
) -> str:
    """
    Create a filesystem-safe filename stem.
    """

    original_stem = Path(filename).stem

    safe_stem = re.sub(
        r"[^A-Za-z0-9_-]+",
        "_",
        original_stem,
    )

    safe_stem = safe_stem.strip(
        "._-"
    )

    if not safe_stem:
        safe_stem = "validated_output"

    return safe_stem[:100]


def ordered_columns(
    *,
    parsed_records: list[dict[str, Any]],
    parser_spec: dict[str, Any],
) -> list[str]:
    """
    Determine output column order.

    First use target-field order from the approved
    parser specification. Then append any additional
    fields found in parsed records.
    """

    mappings = parser_spec.get(
        "mappings",
        [],
    )

    if not isinstance(mappings, list):
        raise ValidatedXlsxError(
            message=(
                "Parser mappings must be a list."
            ),
            guardrail_name=(
                "VALID_PARSER_MAPPINGS_REQUIRED"
            ),
        )

    columns: list[str] = []

    for mapping in mappings:
        if not isinstance(mapping, dict):
            continue

        target = mapping.get("target")

        if (
            isinstance(target, str)
            and target.strip()
            and target not in columns
        ):
            columns.append(
                target
            )

    for record in parsed_records:
        for field_name in record:
            normalized_field_name = str(
                field_name
            )

            if (
                normalized_field_name
                not in columns
            ):
                columns.append(
                    normalized_field_name
                )

    return columns


def calculate_file_sha256(
    file_path: Path,
) -> str:
    """
    Calculate the SHA-256 checksum of a file.
    """

    digest = hashlib.sha256()

    with file_path.open("rb") as file:
        while True:
            block = file.read(
                1024 * 1024
            )

            if not block:
                break

            digest.update(block)

    return digest.hexdigest()


def style_header_row(
    worksheet,
) -> None:
    """
    Apply presentation-ready styling to row 1.
    """

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    header_alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = (
            header_alignment
        )

    worksheet.row_dimensions[1].height = 22


def resize_columns(
    worksheet,
    columns: list[str],
) -> None:
    """
    Set readable column widths using at most the first
    200 data rows to avoid expensive full-sheet scans.
    """

    maximum_inspection_row = min(
        worksheet.max_row,
        201,
    )

    for column_index, column_name in enumerate(
        columns,
        start=1,
    ):
        maximum_length = len(
            str(column_name)
        )

        for row_number in range(
            2,
            maximum_inspection_row + 1,
        ):
            cell_value = worksheet.cell(
                row=row_number,
                column=column_index,
            ).value

            if cell_value is None:
                continue

            maximum_length = max(
                maximum_length,
                len(str(cell_value)),
            )

        column_letter = get_column_letter(
            column_index
        )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            maximum_length + 2,
            50,
        )


def add_validation_summary_sheet(
    *,
    workbook: Workbook,
    run_id: str,
    original_filename: str,
    validation_status: str,
    parsed_record_count: int,
    output_column_count: int,
    validation_summary: dict[str, Any],
    validation_metrics: dict[str, Any],
) -> None:
    """
    Add a summary sheet containing workflow and
    validation information.
    """

    worksheet = workbook.create_sheet(
        "Validation Summary"
    )

    worksheet.append(
        [
            "Field",
            "Value",
        ]
    )

    summary_rows = [
        (
            "Run ID",
            run_id,
        ),
        (
            "Original Filename",
            original_filename,
        ),
        (
            "Validation Status",
            validation_status,
        ),
        (
            "Parsed Row Count",
            parsed_record_count,
        ),
        (
            "Output Column Count",
            output_column_count,
        ),
        (
            "Error Count",
            validation_summary.get(
                "error_count",
                0,
            ),
        ),
        (
            "Warning Count",
            validation_summary.get(
                "warning_count",
                0,
            ),
        ),
        (
            "Row Reconciliation Valid",
            validation_summary.get(
                "row_reconciliation_valid",
                False,
            ),
        ),
        (
            "Completeness Percentage",
            validation_metrics.get(
                "completeness_percentage",
                0,
            ),
        ),
        (
            "Uniqueness Percentage",
            validation_metrics.get(
                "uniqueness_percentage",
                0,
            ),
        ),
        (
            "Valid Record Percentage",
            validation_metrics.get(
                "valid_record_percentage",
                0,
            ),
        ),
        (
            "Duplicate Record Count",
            validation_metrics.get(
                "duplicate_record_count",
                0,
            ),
        ),
        (
            "Missing Required Value Count",
            validation_metrics.get(
                "missing_required_value_count",
                0,
            ),
        ),
        (
            "Invalid Value Count",
            validation_metrics.get(
                "invalid_value_count",
                0,
            ),
        ),
        (
            "Output Truncated",
            validation_metrics.get(
                "output_truncated",
                False,
            ),
        ),
    ]

    for field_name, field_value in (
        summary_rows
    ):
        worksheet.append(
            [
                excel_safe_value(
                    field_name
                ),
                excel_safe_value(
                    field_value
                ),
            ]
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    style_header_row(
        worksheet
    )

    worksheet.column_dimensions[
        "A"
    ].width = 35

    worksheet.column_dimensions[
        "B"
    ].width = 50


def add_validation_issues_sheet(
    *,
    workbook: Workbook,
    validation_issues: list[
        dict[str, Any]
    ],
) -> None:
    """
    Add a validation-issues sheet when warnings or
    sampled validation issues are available.
    """

    if not validation_issues:
        return

    worksheet = workbook.create_sheet(
        "Validation Issues"
    )

    issue_columns = [
        "severity",
        "code",
        "record_number",
        "field",
        "value",
        "message",
    ]

    worksheet.append(
        issue_columns
    )

    for issue in validation_issues:
        if not isinstance(issue, dict):
            continue

        worksheet.append(
            [
                excel_safe_value(
                    issue.get(column)
                )
                for column in issue_columns
            ]
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    style_header_row(
        worksheet
    )

    resize_columns(
        worksheet,
        issue_columns,
    )


def verify_generated_workbook(
    *,
    output_path: Path,
    expected_row_count: int,
    expected_columns: list[str],
) -> None:
    """
    Reopen the generated workbook and verify that it
    contains the expected sheet and dimensions.
    """

    from openpyxl import load_workbook

    try:
        workbook = load_workbook(
            filename=output_path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )

    except Exception as exception:
        raise ValidatedXlsxError(
            message=(
                "Generated XLSX could not be "
                "reopened for verification."
            ),
            guardrail_name=(
                "GENERATED_XLSX_VERIFICATION"
            ),
        ) from exception

    try:
        if (
            "Validated Data"
            not in workbook.sheetnames
        ):
            raise ValidatedXlsxError(
                message=(
                    "Generated XLSX does not contain "
                    "the Validated Data worksheet."
                ),
                guardrail_name=(
                    "VALIDATED_DATA_SHEET_REQUIRED"
                ),
            )

        worksheet = workbook[
            "Validated Data"
        ]

        expected_excel_rows = (
            expected_row_count + 1
        )

        if (
            worksheet.max_row
            != expected_excel_rows
        ):
            raise ValidatedXlsxError(
                message=(
                    "Generated XLSX row count "
                    "does not match parsed output."
                ),
                guardrail_name=(
                    "OUTPUT_ROW_COUNT_VERIFICATION"
                ),
            )

        if (
            worksheet.max_column
            != len(expected_columns)
        ):
            raise ValidatedXlsxError(
                message=(
                    "Generated XLSX column count "
                    "does not match parsed output."
                ),
                guardrail_name=(
                    "OUTPUT_COLUMN_COUNT_VERIFICATION"
                ),
            )

        actual_headers = [
            worksheet.cell(
                row=1,
                column=column_index,
            ).value
            for column_index in range(
                1,
                worksheet.max_column + 1,
            )
        ]

        if (
            actual_headers
            != expected_columns
        ):
            raise ValidatedXlsxError(
                message=(
                    "Generated XLSX headers do not "
                    "match expected output columns."
                ),
                guardrail_name=(
                    "OUTPUT_HEADER_VERIFICATION"
                ),
            )

    finally:
        workbook.close()


def write_validated_xlsx(
    *,
    run_id: str,
    original_filename: str,
    parsed_records: list[
        dict[str, Any]
    ],
    parser_spec: dict[str, Any],
    validation_status: str,
    validation_summary: dict[str, Any],
    validation_metrics: dict[str, Any],
    validation_issues: list[
        dict[str, Any]
    ],
) -> dict[str, Any]:
    """
    Create a new validated XLSX file.

    The original uploaded workbook is never modified.

    Generated worksheets:
    - Validated Data
    - Validation Summary
    - Validation Issues, when issues exist
    """

    if validation_status not in {
        "PASS",
        "WARN",
    }:
        raise ValidatedXlsxError(
            message=(
                "Final XLSX can only be created "
                "after PASS or WARN validation."
            ),
            guardrail_name=(
                "SUCCESSFUL_VALIDATION_REQUIRED"
            ),
        )

    if not isinstance(
        parsed_records,
        list,
    ):
        raise ValidatedXlsxError(
            message=(
                "Parsed records must be a list."
            ),
            guardrail_name=(
                "PARSED_RECORDS_REQUIRED"
            ),
        )

    if not isinstance(
        parser_spec,
        dict,
    ):
        raise ValidatedXlsxError(
            message=(
                "Approved parser specification "
                "must be a dictionary."
            ),
            guardrail_name=(
                "APPROVED_PARSER_SPEC_REQUIRED"
            ),
        )

    if not isinstance(
        validation_summary,
        dict,
    ):
        raise ValidatedXlsxError(
            message=(
                "Validation summary must "
                "be a dictionary."
            ),
            guardrail_name=(
                "VALIDATION_SUMMARY_REQUIRED"
            ),
        )

    if not isinstance(
        validation_metrics,
        dict,
    ):
        raise ValidatedXlsxError(
            message=(
                "Validation metrics must "
                "be a dictionary."
            ),
            guardrail_name=(
                "VALIDATION_METRICS_REQUIRED"
            ),
        )

    if not isinstance(
        validation_issues,
        list,
    ):
        raise ValidatedXlsxError(
            message=(
                "Validation issues must be a list."
            ),
            guardrail_name=(
                "VALIDATION_ISSUES_REQUIRED"
            ),
        )

    for record_index, record in enumerate(
        parsed_records,
        start=1,
    ):
        if not isinstance(record, dict):
            raise ValidatedXlsxError(
                message=(
                    "Parsed record at position "
                    f"{record_index} is not "
                    "a dictionary."
                ),
                guardrail_name=(
                    "DICTIONARY_RECORD_REQUIRED"
                ),
            )

    columns = ordered_columns(
        parsed_records=parsed_records,
        parser_spec=parser_spec,
    )

    if not columns:
        raise ValidatedXlsxError(
            message=(
                "At least one output column "
                "is required."
            ),
            guardrail_name=(
                "OUTPUT_COLUMNS_REQUIRED"
            ),
        )

    output_directory = Path(
        settings.output_directory
    )

    output_directory.mkdir(
        parents=True,
        exist_ok=True,
    )

    safe_stem = (
        sanitize_output_filename_stem(
            original_filename
        )
    )

    output_filename = (
        f"{safe_stem}_validated_"
        f"{run_id[:8]}.xlsx"
    )

    output_path = (
        output_directory
        / output_filename
    )

    if output_path.exists():
        raise ValidatedXlsxError(
            message=(
                "Validated XLSX output "
                "already exists."
            ),
            guardrail_name=(
                "IMMUTABLE_OUTPUT_PATH"
            ),
        )

    workbook = Workbook()

    try:
        data_sheet = workbook.active
        data_sheet.title = (
            "Validated Data"
        )

        data_sheet.append(
            columns
        )

        for record in parsed_records:
            data_sheet.append(
                [
                    excel_safe_value(
                        record.get(column)
                    )
                    for column in columns
                ]
            )

        data_sheet.freeze_panes = "A2"
        data_sheet.auto_filter.ref = (
            data_sheet.dimensions
        )

        style_header_row(
            data_sheet
        )

        resize_columns(
            data_sheet,
            columns,
        )

        add_validation_summary_sheet(
            workbook=workbook,
            run_id=run_id,
            original_filename=(
                original_filename
            ),
            validation_status=(
                validation_status
            ),
            parsed_record_count=len(
                parsed_records
            ),
            output_column_count=len(
                columns
            ),
            validation_summary=(
                validation_summary
            ),
            validation_metrics=(
                validation_metrics
            ),
        )

        add_validation_issues_sheet(
            workbook=workbook,
            validation_issues=(
                validation_issues
            ),
        )

        workbook.save(
            output_path
        )

    except Exception:
        output_path.unlink(
            missing_ok=True
        )
        raise

    finally:
        workbook.close()

    if not output_path.exists():
        raise ValidatedXlsxError(
            message=(
                "Validated XLSX was not created."
            ),
            guardrail_name=(
                "OUTPUT_FILE_CREATION_REQUIRED"
            ),
        )

    output_size_bytes = (
        output_path.stat().st_size
    )

    if output_size_bytes <= 0:
        output_path.unlink(
            missing_ok=True
        )

        raise ValidatedXlsxError(
            message=(
                "Validated XLSX is empty."
            ),
            guardrail_name=(
                "NON_EMPTY_OUTPUT_FILE"
            ),
        )

    try:
        verify_generated_workbook(
            output_path=output_path,
            expected_row_count=len(
                parsed_records
            ),
            expected_columns=columns,
        )

    except Exception:
        output_path.unlink(
            missing_ok=True
        )
        raise

    output_checksum = (
        calculate_file_sha256(
            output_path
        )
    )

    worksheet_names = [
        "Validated Data",
        "Validation Summary",
    ]

    if validation_issues:
        worksheet_names.append(
            "Validation Issues"
        )

    return {
        "final_output_path": str(
            output_path
        ),
        "final_output_filename": (
            output_filename
        ),
        "final_output_size_bytes": (
            output_size_bytes
        ),
        "final_output_row_count": len(
            parsed_records
        ),
        "final_output_column_count": len(
            columns
        ),
        "final_output_checksum_sha256": (
            output_checksum
        ),
        "final_output_available": True,
        "worksheet_names": (
            worksheet_names
        ),
    }
