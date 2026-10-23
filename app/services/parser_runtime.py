import math
import re
from datetime import date, datetime
from decimal import (
    Decimal,
    InvalidOperation,
)
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.core.config import settings
from app.services.parser_registry import (
    validate_parser_specification,
)


class ParserExecutionError(ValueError):
    def __init__(
        self,
        message: str,
        guardrail_name: str,
        *,
        row_number: int | None = None,
        column_name: str | None = None,
    ) -> None:
        super().__init__(message)

        self.message = message
        self.guardrail_name = guardrail_name
        self.row_number = row_number
        self.column_name = column_name


def normalize_column_name(
    value: Any,
) -> str:
    if value is None:
        return ""

    normalized = re.sub(
        r"[^a-z0-9]+",
        "_",
        str(value).strip().lower(),
    )

    return normalized.strip("_")


def make_unique_headers(
    headers: list[str],
) -> list[str]:
    counts: dict[str, int] = {}
    unique_headers: list[str] = []

    for position, header in enumerate(
        headers,
        start=1,
    ):
        base_name = (
            header
            if header
            else f"unnamed_column_{position}"
        )

        counts[base_name] = (
            counts.get(base_name, 0) + 1
        )

        if counts[base_name] == 1:
            unique_name = base_name
        else:
            unique_name = (
                f"{base_name}_"
                f"{counts[base_name]}"
            )

        unique_headers.append(unique_name)

    return unique_headers


def is_empty_value(
    value: Any,
) -> bool:
    if value is None:
        return True

    if (
        isinstance(value, str)
        and not value.strip()
    ):
        return True

    return False


def apply_identity(
    value: Any,
) -> Any:
    return value


def apply_trim(
    value: Any,
) -> Any:
    if value is None:
        return None

    if isinstance(value, str):
        return value.strip()

    return value


def apply_upper(
    value: Any,
) -> Any:
    if value is None:
        return None

    return str(value).strip().upper()


def apply_lower(
    value: Any,
) -> Any:
    if value is None:
        return None

    return str(value).strip().lower()


def apply_parse_date(
    value: Any,
) -> str | None:
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date().isoformat()

    if isinstance(value, date):
        return value.isoformat()

    text = str(value).strip()

    accepted_formats = (
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%m/%d/%Y",
    )

    for date_format in accepted_formats:
        try:
            parsed_value = datetime.strptime(
                text,
                date_format,
            )

            return parsed_value.date().isoformat()

        except ValueError:
            continue

    raise ParserExecutionError(
        message=(
            f"Value '{text}' cannot be parsed "
            "as an accepted date."
        ),
        guardrail_name=(
            "VALID_DATE_REQUIRED"
        ),
    )


def apply_parse_decimal(
    value: Any,
) -> str | None:
    if value is None:
        return None

    if isinstance(value, bool):
        raise ParserExecutionError(
            message=(
                "Boolean values cannot be parsed "
                "as decimals."
            ),
            guardrail_name=(
                "VALID_DECIMAL_REQUIRED"
            ),
        )

    if isinstance(value, float):
        if not math.isfinite(value):
            raise ParserExecutionError(
                message=(
                    "Non-finite decimal values "
                    "are not permitted."
                ),
                guardrail_name=(
                    "FINITE_DECIMAL_REQUIRED"
                ),
            )

    normalized_value = value

    if isinstance(value, str):
        normalized_value = (
            value.strip()
            .replace(",", "")
        )

        if not normalized_value:
            return None

    try:
        decimal_value = Decimal(
            str(normalized_value)
        )

    except (
        InvalidOperation,
        ValueError,
    ) as exception:
        raise ParserExecutionError(
            message=(
                f"Value '{value}' cannot be "
                "parsed as a decimal."
            ),
            guardrail_name=(
                "VALID_DECIMAL_REQUIRED"
            ),
        ) from exception

    if not decimal_value.is_finite():
        raise ParserExecutionError(
            message=(
                "Non-finite decimal values "
                "are not permitted."
            ),
            guardrail_name=(
                "FINITE_DECIMAL_REQUIRED"
            ),
        )

    return format(
        decimal_value,
        "f",
    )


TRANSFORMERS = {
    "identity": apply_identity,
    "trim": apply_trim,
    "upper": apply_upper,
    "lower": apply_lower,
    "parse_date": apply_parse_date,
    "parse_decimal": apply_parse_decimal,
}


def validate_execution_authorization(
    *,
    parser_execution_authorized: bool,
    parser_match_found: bool,
    matched_parser_approval_status: str | None,
    matched_parser_status: str | None,
) -> None:
    if parser_execution_authorized is not True:
        raise ParserExecutionError(
            message=(
                "Parser execution is not authorized."
            ),
            guardrail_name=(
                "PARSER_EXECUTION_AUTHORIZATION_REQUIRED"
            ),
        )

    if parser_match_found is not True:
        raise ParserExecutionError(
            message=(
                "An exact parser match is required "
                "before execution."
            ),
            guardrail_name=(
                "EXACT_PARSER_MATCH_REQUIRED"
            ),
        )

    if (
        matched_parser_approval_status
        != "APPROVED"
    ):
        raise ParserExecutionError(
            message=(
                "The matched parser version "
                "is not approved."
            ),
            guardrail_name=(
                "APPROVED_PARSER_REQUIRED"
            ),
        )

    if matched_parser_status != "ACTIVE":
        raise ParserExecutionError(
            message=(
                "The matched parser version "
                "is not active."
            ),
            guardrail_name=(
                "ACTIVE_PARSER_VERSION_REQUIRED"
            ),
        )


def transform_value(
    *,
    value: Any,
    transformation: str,
    row_number: int,
    source_column: str,
) -> Any:
    transformer = TRANSFORMERS.get(
        transformation
    )

    if transformer is None:
        raise ParserExecutionError(
            message=(
                f"Transformation '{transformation}' "
                "is not supported."
            ),
            guardrail_name=(
                "ALLOWLISTED_TRANSFORMATION_ONLY"
            ),
            row_number=row_number,
            column_name=source_column,
        )

    try:
        return transformer(value)

    except ParserExecutionError as exception:
        if exception.row_number is None:
            exception.row_number = row_number

        if exception.column_name is None:
            exception.column_name = (
                source_column
            )

        raise


def execute_declarative_parser(
    *,
    stored_path: str,
    parser_spec: dict[str, Any],
    parser_execution_authorized: bool,
    parser_match_found: bool,
    matched_parser_approval_status: str | None,
    matched_parser_status: str | None,
) -> dict[str, Any]:
    """
    Execute an approved declarative parser.

    This function does not execute generated code.
    """

    validate_execution_authorization(
        parser_execution_authorized=(
            parser_execution_authorized
        ),
        parser_match_found=(
            parser_match_found
        ),
        matched_parser_approval_status=(
            matched_parser_approval_status
        ),
        matched_parser_status=(
            matched_parser_status
        ),
    )

    validate_parser_specification(
        parser_spec
    )

    mappings = parser_spec["mappings"]

    if (
        len(mappings)
        > settings.maximum_parser_mappings
    ):
        raise ParserExecutionError(
            message=(
                "Parser mapping count exceeds "
                "the configured limit of "
                f"{settings.maximum_parser_mappings}."
            ),
            guardrail_name=(
                "MAXIMUM_PARSER_MAPPINGS"
            ),
        )

    file_path = Path(stored_path)

    if (
        not file_path.exists()
        or not file_path.is_file()
    ):
        raise ParserExecutionError(
            message=(
                "Stored workbook is unavailable."
            ),
            guardrail_name=(
                "STORED_FILE_REQUIRED"
            ),
        )

    try:
        workbook = load_workbook(
            filename=file_path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )

    except Exception as exception:
        raise ParserExecutionError(
            message=(
                "Stored workbook could not be opened "
                "for parser execution."
            ),
            guardrail_name=(
                "WORKBOOK_OPEN_VALIDATION"
            ),
        ) from exception

    try:
        source_sheet = parser_spec[
            "source_sheet"
        ]

        if source_sheet not in workbook.sheetnames:
            raise ParserExecutionError(
                message=(
                    f"Configured source sheet "
                    f"'{source_sheet}' does not exist."
                ),
                guardrail_name=(
                    "SOURCE_SHEET_EXISTS"
                ),
            )

        worksheet = workbook[source_sheet]

        header_row = parser_spec[
            "header_row"
        ]

        raw_headers = [
            worksheet.cell(
                row=header_row,
                column=column_number ).value
            for column_number in range(
                1,
                worksheet.max_column + 1,
            )
        ]

        normalized_headers = (
            make_unique_headers(
                [
                    normalize_column_name(value)
                    for value in raw_headers
                ]
            )
        )

        header_positions = {
            header: position
            for position, header in enumerate(
                normalized_headers,
                start=1,
            )
        }

        required_source_columns = {
            mapping["source"]
            for mapping in mappings
        }

        missing_source_columns = sorted(
            required_source_columns
            - set(header_positions)
        )

        if missing_source_columns:
            raise ParserExecutionError(
                message=(
                    "Configured source columns "
                    "were not found: "
                    f"{missing_source_columns}"
                ),
                guardrail_name=(
                    "SOURCE_COLUMNS_EXIST"
                ),
            )

        parsed_records: list[
            dict[str, Any]
        ] = []

        source_row_count = max(
            worksheet.max_row - header_row,
            0,
        )

        rejected_row_count = 0

        data_start_row = header_row + 1

        maximum_source_row = min(
            worksheet.max_row,
            header_row
            + settings.maximum_parser_output_rows,
        )

        for row_number in range(
            data_start_row,
            maximum_source_row + 1,
        ):
            row_values = [
                worksheet.cell(
                    row=row_number,
                    column=column_number,
                ).value
                for column_number in range(
                    1,
                    worksheet.max_column + 1,
                )
            ]

            if all(
                is_empty_value(value)
                for value in row_values
            ):
                continue

            parsed_record: dict[
                str,
                Any,
            ] = {}

            try:
                for mapping in mappings:
                    source_column = mapping[
                        "source"
                    ]

                    source_position = (
                        header_positions[
                            source_column
                        ]
                    )

                    source_value = (
                        worksheet.cell(
                            row=row_number,
                            column=source_position,
                        ).value
                    )

                    if (
                        mapping["required"]
                        and is_empty_value(
                            source_value
                        )
                    ):
                        raise ParserExecutionError(
                            message=(
                                "Required source value "
                                "is empty."
                            ),
                            guardrail_name=(
                                "REQUIRED_VALUE_PRESENT"
                            ),
                            row_number=row_number,
                            column_name=source_column,
                        )

                    transformed_value = (
                        transform_value(
                            value=source_value,
                            transformation=(
                                mapping["transform"]
                            ),
                            row_number=row_number,
                            source_column=(
                                source_column
                            ),
                        )
                    )

                    parsed_record[
                        mapping["target"]
                    ] = transformed_value

                parsed_records.append(
                    parsed_record
                )

            except ParserExecutionError:
                rejected_row_count += 1
                raise

        output_truncated = (
            source_row_count
            > settings.maximum_parser_output_rows
        )

        return {
            "parsed_records": parsed_records,
            "parsed_row_count": len(
                parsed_records
            ),
            "source_row_count": (
                source_row_count
            ),
            "rejected_row_count": (
                rejected_row_count
            ),
            "output_truncated": (
                output_truncated
            ),
            "parser_execution_summary": {
                "source_sheet": (
                    source_sheet
                ),
                "header_row": header_row,
                "mapping_count": len(mappings),
                "maximum_output_rows": (
                    settings
                    .maximum_parser_output_rows
                ),
                "transformations_used": sorted(
                    {
                        mapping["transform"]
                        for mapping in mappings
                    }
                ),
                "execution_mode": (
                    "DECLARATIVE_ALLOWLISTED"
                ),
                "dynamic_code_execution": False,
            },
        }

    finally:
        workbook.close()
