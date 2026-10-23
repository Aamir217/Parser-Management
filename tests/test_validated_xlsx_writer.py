from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.services.validated_xlsx_writer import (
    ValidatedXlsxError,
    excel_safe_value,
    write_validated_xlsx,
)


def parser_spec() -> dict:
    return {
        "source_sheet": "Trade Data",
        "header_row": 1,
        "mappings": [
            {
                "source": "trade_id",
                "target": "trade_id",
                "transform": "trim",
                "required": True,
            },
            {
                "source": "amount",
                "target": "amount",
                "transform": "parse_decimal",
                "required": True,
            },
        ],
    }


def test_formula_injection_is_escaped():
    assert (
        excel_safe_value("=2+2")
        == "'=2+2"
    )

    assert (
        excel_safe_value("@SUM(A1:A2)")
        == "'@SUM(A1:A2)"
    )


def test_validated_xlsx_is_created(
    monkeypatch,
    tmp_path,
):
    monkeypatch.setattr(
        (
            "app.services."
            "validated_xlsx_writer."
            "settings.output_directory"
        ),
        str(tmp_path),
    )

    run_id = str(uuid4())

    result = write_validated_xlsx(
        run_id=run_id,
        original_filename=(
            "trades.xlsx"
        ),
        parsed_records=[
            {
                "trade_id": "TRD-1001",
                "amount": "125000.50",
            },
            {
                "trade_id": "TRD-1002",
                "amount": "98000",
            },
        ],
        parser_spec=parser_spec(),
        validation_status="PASS",
        validation_summary={
            "error_count": 0,
            "warning_count": 0,
        },
        validation_metrics={
            "completeness_percentage": 100,
            "uniqueness_percentage": 100,
            "valid_record_percentage": 100,
        },
        validation_issues=[],
    )

    output_path = Path(
        result["final_output_path"]
    )

    assert output_path.exists()

    assert (
        result["final_output_available"]
        is True
    )

    assert (
        result["final_output_row_count"]
        == 2
    )

    assert (
        result[
            "final_output_column_count"
        ]
        == 2
    )

    assert len(
        result[
            "final_output_checksum_sha256"
        ]
    ) == 64

    workbook = load_workbook(
        output_path,
        read_only=True,
        data_only=True,
    )

    assert (
        "Validated Data"
        in workbook.sheetnames
    )

    assert (
        "Validation Summary"
        in workbook.sheetnames
    )

    data_sheet = workbook[
        "Validated Data"
    ]

    assert data_sheet["A1"].value == (
        "trade_id"
    )

    assert data_sheet["B1"].value == (
        "amount"
    )

    assert data_sheet["A2"].value == (
        "TRD-1001"
    )

    workbook.close()


def test_failed_validation_cannot_export(
    monkeypatch,
    tmp_path,
):
    monkeypatch.setattr(
        (
            "app.services."
            "validated_xlsx_writer."
            "settings.output_directory"
        ),
        str(tmp_path),
    )

    with pytest.raises(
        ValidatedXlsxError
    ) as exception:
        write_validated_xlsx(
            run_id=str(uuid4()),
            original_filename=(
                "failed.xlsx"
            ),
            parsed_records=[],
            parser_spec=parser_spec(),
            validation_status="FAIL",
            validation_summary={},
            validation_metrics={},
            validation_issues=[],
        )

    assert (
        exception.value.guardrail_name
        == (
            "SUCCESSFUL_VALIDATION_REQUIRED"
        )
    )
