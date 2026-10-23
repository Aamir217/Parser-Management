from pathlib import Path

from openpyxl import Workbook

from app.services.workbook_inspector import (
    inspect_workbook_file,
    normalize_name,
)


def create_inspection_workbook(
    file_path: Path,
) -> None:
    workbook = Workbook()

    trades = workbook.active
    trades.title = "Trade Data"

    trades.append([
        "Trade ID",
        "Trade Date",
        "Currency",
        "Amount",
        "Status",
    ])

    trades.append([
        "TRD-1001",
        "2026-10-19",
        "USD",
        1000.50,
        "BOOKED",
    ])

    trades.append([
        "TRD-1002",
        "2026-10-20",
        "EUR",
        2000,
        "PENDING",
    ])

    trades["F1"] = "Calculated Amount"
    trades["F2"] = "=D2*2"

    reference = workbook.create_sheet(
        "Reference Data"
    )

    reference.sheet_state = "hidden"

    reference.append([
        "Currency Code",
        "Currency Name",
    ])

    reference.append([
        "USD",
        "US Dollar",
    ])

    reference.merge_cells(
        "A4:B4"
    )

    reference["A4"] = "Merged title"

    workbook.save(file_path)


def test_normalize_name():
    assert normalize_name(
        "Trade Date"
    ) == "trade_date"

    assert normalize_name(
        " Counterparty-ID "
    ) == "counterparty_id"


def test_workbook_inspection(tmp_path):
    workbook_path = (
        tmp_path
        / "inspection_test.xlsx"
    )

    create_inspection_workbook(
        workbook_path
    )

    result = inspect_workbook_file(
        str(workbook_path)
    )

    assert result["inspection_version"] == (
        "1.0"
    )

    assert result["sheet_count"] == 2

    assert result[
        "hidden_sheet_count"
    ] == 1

    assert result[
        "total_formula_count"
    ] == 1

    assert result[
        "total_merged_range_count"
    ] == 1

    trades = result["sheets"][0]

    assert trades["name"] == "Trade Data"

    assert trades["normalized_name"] == (
        "trade_data"
    )

    assert trades["header_row"] == 1

    assert (
        "trade_id"
        in trades["normalized_headers"]
    )

    assert (
        "trade_date"
        in trades["normalized_headers"]
    )

    assert trades["formula_count"] == 1

    assert "F2" in trades[
        "formula_cells"
    ]

    reference = result["sheets"][1]

    assert reference["is_hidden"] is True

    assert reference[
        "merged_range_count"
    ] == 1

    assert "A4:B4" in reference[
        "merged_ranges"
    ]


def test_column_type_inference(tmp_path):
    workbook_path = (
        tmp_path
        / "type_test.xlsx"
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"

    sheet.append([
        "Identifier",
        "Amount",
        "Active",
    ])

    sheet.append([
        "ID-1",
        100.50,
        True,
    ])

    sheet.append([
        "ID-2",
        200,
        False,
    ])

    workbook.save(workbook_path)

    result = inspect_workbook_file(
        str(workbook_path)
    )

    profiles = {
        profile["normalized_name"]:
        profile
        for profile
        in result["sheets"][0][
            "column_profiles"
        ]
    }

    assert profiles[
        "identifier"
    ]["inferred_type"] == "string"

    assert profiles[
        "active"
    ]["inferred_type"] == "boolean"

    assert profiles[
        "amount"
    ]["inferred_type"] in {
        "integer",
        "decimal",
    }
